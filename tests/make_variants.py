# Generuje trudne warianty jadłospisu (październik 2026) do testów parsera
import os
from docx import Document

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'files')
SOUPS = ['Zupa pomidorowa z ryżem', 'Krupnik z ziemniakami', 'Rosół z makaronem', 'Barszcz czerwony', 'Ogórkowa z koperkiem']
MAINS = ['Kotlet schabowy, ziemniaki, surówka', 'Spaghetti bolognese', 'Pierogi ruskie ze śmietaną',
         'Filet z miruny, puree, marchewka', 'Naleśniki z serem']


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    return t


# 1. Inna kolejność kolumn, synonimy i literówki w nagłówkach, różne formaty dat,
#    literówki w nazwach dni, dzień wolny w scalonej komórce, zła nazwa dnia przy dacie
doc = Document()
doc.add_paragraph('Jadłospis październik 2026')
rows = [['Data', 'II danie', 'Zupa dnia', 'Alergny', 'Deser'],
        ['Poniedzialek 5.10', MAINS[0], SOUPS[0], '1,7', ''],
        ['Wtorke 6/10', MAINS[1], SOUPS[1], '1,3', 'jabłko'],
        ['Środa 7 października', MAINS[2], SOUPS[2], '1,7,9', ''],
        ['Czwartek 8.X', MAINS[3], SOUPS[3], '4', ''],
        ['Piatek 9.10.2026', MAINS[4], SOUPS[4], '1,3,7', 'kisiel'],
        ['Pon. 12.10', MAINS[1], SOUPS[0], '', ''],
        ['Wt 13.10', MAINS[2], SOUPS[1], '', ''],
        ['Śr 14.10', 'DZIEŃ WOLNY - Dzień Edukacji Narodowej', '', '', ''],
        ['Czwartek 16.10', MAINS[3], SOUPS[3], '', '']]
t = table(doc, rows)
t.cell(8, 1).merge(t.cell(8, 4))
t.cell(8, 1).text = 'DZIEŃ WOLNY - Dzień Edukacji Narodowej'
doc.save(os.path.join(OUT, 'v1_kolumny_literowki.docx'))

# 2. Same nazwy dni bez dat (daty trzeba wyliczyć z miesiąca w tytule)
doc = Document()
doc.add_paragraph('JADŁOSPIS – PAŹDZIERNIK 2026')
weeks = [['Czwartek', 'Piątek'], ['Poniedziałek', 'Wtorek', 'Środa', 'Czwartek', 'Piątek'],
         ['Poniedziałek', 'Wtorek', 'Środa', 'Czwartek', 'Piątek']]
for n, week in enumerate(weeks, 1):
    doc.add_paragraph(f'Tydzień {n}')
    table(doc, [['Dzień', 'Zupa', 'Drugie danie']] + [[d, SOUPS[i % 5], MAINS[i % 5]] for i, d in enumerate(week)])
doc.save(os.path.join(OUT, 'v2_bez_dat.docx'))

# 3. Dni w kolumnach, rodzaje dań w wierszach
doc = Document()
doc.add_paragraph('Jadłospis na tydzień 5-9 października 2026')
table(doc, [['', 'Pon 5.10', 'Wt 6.10', 'Śr 7.10', 'Czw 8.10', 'Pt 9.10'],
            ['Zupa'] + SOUPS,
            ['Drugie danie'] + MAINS,
            ['Deser', '', 'owoc', '', '', 'galaretka'],
            ['Alergeny', '1', '1,3', '7', '4', '1,7']])
doc.save(os.path.join(OUT, 'v3_dni_w_kolumnach.docx'))

# 4. Punkty (lista wypunktowana) z etykietami
doc = Document()
doc.add_paragraph('Jadłospis – październik 2026')
for i, (d, n) in enumerate([('Poniedziałek', 5), ('Wtorek', 6), ('Środa', 7), ('Czwartek', 8), ('Piątek', 9)]):
    doc.add_paragraph(f'{d}, {n} października', style='List Bullet')
    doc.add_paragraph(f'Zupa: {SOUPS[i].lower()}', style='List Bullet 2')
    doc.add_paragraph(f'II danie – {MAINS[i]}', style='List Bullet 2')
    if i % 2 == 0:
        doc.add_paragraph('Deser: jabłko', style='List Bullet 2')
    doc.add_paragraph('Alergeny: 1, 7', style='List Bullet 2')
doc.save(os.path.join(OUT, 'v4_punkty.docx'))

# 5. Zwykłe akapity bez etykiet, z legendą alergenów i cennikiem na końcu
doc = Document()
doc.add_paragraph('JADŁOSPIS SZKOLNY')
for i, (d, n) in enumerate([('PONIEDZIAŁEK', 5), ('WTOREK', 6), ('ŚRODA', 7), ('CZWARTEK', 8), ('PIĄTEK', 9)]):
    doc.add_paragraph(f'{d} {n:02d}.10.2026')
    doc.add_paragraph(SOUPS[i])
    doc.add_paragraph(MAINS[i])
doc.add_paragraph('ALERGENY – LEGENDA')
doc.add_paragraph('1 – zboża zawierające gluten')
doc.add_paragraph('7 – mleko i produkty mleczne')
doc.add_paragraph('CENNIK')
doc.add_paragraph('Zupa – 8,50 zł')
doc.save(os.path.join(OUT, 'v5_akapity.docx'))

# 6. Scalone komórki: deser wspólny dla całego tygodnia (pionowo)
doc = Document()
doc.add_paragraph('JADŁOSPIS PAŹDZIERNIK 2026')
rows = [['DZIEŃ', 'ZUPA', 'DRUGIE DANIE', 'DESER']] + \
       [[f'{d} {n}.10', SOUPS[i], MAINS[i], 'Owoc sezonowy' if i == 0 else '']
        for i, (d, n) in enumerate([('PONIEDZIAŁEK', 19), ('WTOREK', 20), ('ŚRODA', 21), ('CZWARTEK', 22), ('PIĄTEK', 23)])]
t = table(doc, rows)
merged = t.cell(1, 3).merge(t.cell(5, 3))
merged.text = 'Owoc sezonowy'
doc.save(os.path.join(OUT, 'v6_scalone.docx'))
print('gotowe:', sorted(f for f in os.listdir(OUT) if f.startswith('v')))
